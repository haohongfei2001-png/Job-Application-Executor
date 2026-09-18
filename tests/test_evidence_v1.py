from docx import Document

from executor.evidence import ProfileBuilder


def test_max_docx_builds_provenance_and_keeps_pending_unset(tmp_path):
    path = tmp_path / "max.docx"
    doc = Document()
    doc.add_paragraph("姓名：Example User")
    doc.add_paragraph("现居住城市：测试市（2026-09-17确认）")
    doc.add_paragraph("驾驶证：C2机动车驾驶证")
    doc.add_paragraph("驾驶证：C2")
    doc.add_paragraph("生源地：[待补充，如网申要求]")
    doc.save(path)

    builder = ProfileBuilder().import_max_docx(path)
    profile = builder.build()

    assert profile.fields["identity.full_name"].value == "Example User"
    city = profile.fields["identity.current_city"]
    assert city.user_confirmed is True
    assert city.last_verified == "2026-09-17"
    assert city.sources[0].kind == "max_docx"
    assert "identity.student_origin" not in profile.fields
    assert profile.fields["identity.driver_license"].value == "C2"
    assert not profile.collections.get("conflicts")


def test_project_merge_normalizes_whitespace():
    from executor.evidence import _merge_projects

    merged = _merge_projects(
        [{"title": "PQCD 框架下 B→K", "metadata": [], "bullets": [], "source_kind": "max_docx"}],
        [{"title": "PQCD框架下 B→K", "metadata": [], "bullets": ["result"], "source_kind": "resume_pdf"}],
    )
    assert len(merged) == 1
    assert merged[0]["bullets"] == ["result"]
    assert set(merged[0]["sources"]) == {"max_docx", "resume_pdf"}


def test_verified_history_requires_explicit_canonical_mapping(tmp_path):
    import json

    raw = tmp_path / "history.json"
    raw.write_text(json.dumps({
        "site_question_123": "old raw answer",
        "canonical_fields": {
            "identity.current_city": {
                "value": "Verified City",
                "verified_at": "2026-01-01",
                "confidence": 0.92,
            }
        },
    }), encoding="utf-8")
    profile = ProfileBuilder().import_historical_json("fixture", raw).build()
    assert profile.fields["identity.current_city"].value == "Verified City"
    assert profile.fields["identity.current_city"].sources[0].kind == "verified_history"
    assert "site_question_123" not in profile.fields


def test_user_explicit_override_survives_profile_rebuild_and_outranks_confirmed_max(tmp_path):
    from executor.evidence import set_user_confirmed_field, write_profile

    source = tmp_path / "max.docx"
    doc = Document()
    doc.add_paragraph("现居住城市：Document City（2026-09-17确认）")
    doc.save(source)

    profile_path = tmp_path / "profile.json"
    write_profile(ProfileBuilder().import_max_docx(source).build(), profile_path)
    set_user_confirmed_field(
        profile_path,
        "identity.current_city",
        "Explicit City",
        note="test confirmation",
    )

    rebuilt = (
        ProfileBuilder()
        .import_max_docx(source)
        .import_user_overrides(profile_path)
        .build()
    )
    field = rebuilt.fields["identity.current_city"]
    assert field.value == "Explicit City"
    assert field.user_confirmed is True
    assert any(x.kind == "user_explicit" for x in field.sources)

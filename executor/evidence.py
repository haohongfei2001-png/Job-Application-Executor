from __future__ import annotations

import hashlib
import fcntl
import json
import os
import re
import tempfile
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document

from .models import ApplicantProfile, AssetRef, EvidenceRef, ProfileField
from .profile import aliases_for, is_sensitive_key


LABEL_MAP = {
    "姓名": "identity.full_name",
    "英文名": "identity.english_name",
    "性别": "identity.gender",
    "出生日期": "identity.birth_date",
    "民族": "identity.ethnicity",
    "政治面貌": "identity.political_status",
    "证件类型": "identity.id_type",
    "身份证号": "identity.id_number",
    "证件号码": "identity.id_number",
    "手机号码": "identity.phone",
    "手机号": "identity.phone",
    "邮箱": "identity.email",
    "籍贯": "identity.native_place",
    "户籍所在地": "identity.domicile",
    "户籍": "identity.domicile",
    "家庭住址": "identity.family_address",
    "家庭住址（网申复制）": "identity.family_address",
    "高考前户口所在地": "identity.pre_gaokao_domicile",
    "现居住地": "identity.current_residence",
    "现居住城市": "identity.current_city",
    "通信地址": "identity.mailing_address",
    "紧急联系人姓名": "identity.emergency_contact_name",
    "紧急联系人联系方式": "identity.emergency_contact_phone",
    "婚姻状况": "identity.marital_status",
    "驾驶证": "identity.driver_license",
    "生源地": "identity.student_origin",
    "户口类别": "identity.household_type",
    "健康状况": "identity.health_status",
    "人事档案所在单位": "identity.personnel_file_place",
    "是否具有外国国籍或境外长期居留身份": "identity.foreign_residency_status",
    "是否具有外国国籍，或者国（境）外永久居留权、长期居留许可等境外身份": "identity.foreign_residency_status",
    "中国东方员工工作回避要求": "compliance.coamc_employee_recusal_requirements_met",
    "是否符合中国东方员工工作回避有关要求": "compliance.coamc_employee_recusal_requirements_met",
    "家庭成员1姓名": "family.primary.name",
    "家庭成员1关系": "family.primary.relationship",
    "家庭成员1工作单位": "family.primary.work_unit",
    "家庭成员1部门及职务": "family.primary.department_title",
    "家庭成员1工作所在地": "family.primary.work_location",
    "网申隐私政策自动决策": "policy.auto_accept_privacy_terms",
    "网申真实性/投递声明自动决策": "policy.auto_accept_truth_submission_declarations",
    "新公司法律/合规声明自动决策": "policy.auto_decide_company_legal_compliance",
    "最终投递必须本人点击": "policy.final_submission_requires_user_click",
    "求职身份": "preferences.candidate_status",
    "预计毕业日期": "education.highest.graduation_date",
    "培养/学历取得方式": "education.highest.study_type",
    "期望工作城市": "preferences.preferred_cities",
    "是否接受工作地点调剂": "preferences.accept_location_adjustment",
    "是否接受岗位调剂": "preferences.accept_role_adjustment",
    "是否接受出差": "preferences.accept_travel",
    "是否接受长期驻外": "preferences.accept_long_term_assignment",
    "最早到岗时间": "preferences.earliest_start_date",
    "期望薪资": "preferences.salary_policy",
    "最高学历": "education.highest.degree",
    "最高学历学校": "education.highest.school",
    "最高学历院系": "education.highest.college",
    "最高学历专业": "education.highest.major",
    "最高学历入学时间": "education.highest.start_date",
    "最高学历毕业时间": "education.highest.graduation_date",
    "最高学历受教育类型": "education.highest.study_type",
    "最高学历学习成绩排名档位": "education.highest.rank_band",
    "本科院校": "education.bachelor.school",
    "本科院系": "education.bachelor.college",
    "本科专业": "education.bachelor.major",
    "本科入学时间": "education.bachelor.start_date",
    "本科毕业时间": "education.bachelor.graduation_date",
    "本科受教育类型": "education.bachelor.study_type",
    "本科GPA": "education.bachelor.gpa",
    "培养项目": "education.bachelor.training_program",
    "外语等级": "language.cet6.level",
    "CET-6成绩": "language.cet6.score",
    "普通话": "certificates.mandarin",
    "计算机证书": "certificates.computer",
    "学生干部经历": "experience.student_leadership",
    "校园组织经历": "experience.campus_organization",
    "正式实习经历": "experience.internship",
    "正式工作经历": "experience.work",
    "奖学金 / 荣誉": "awards.summary",
    "个人特长": "narrative.strengths",
    "兴趣爱好": "narrative.interests",
    "自我评价": "narrative.self_evaluation",
}

BOOL_KEYS = {
    "preferences.accept_location_adjustment",
    "preferences.accept_role_adjustment",
    "preferences.accept_travel",
    "preferences.accept_long_term_assignment",
    "identity.foreign_residency_status",
    "compliance.coamc_employee_recusal_requirements_met",
    "policy.auto_accept_privacy_terms",
    "policy.auto_accept_truth_submission_declarations",
    "policy.auto_decide_company_legal_compliance",
    "policy.final_submission_requires_user_click",
}

PENDING_MARKERS = ("待补充", "待查", "未固定", "根据毕业", "建议按")
CONFIRM_RE = re.compile(r"(20\d{2}-\d{2}-\d{2})确认")
ANNOTATION_RE = re.compile(r"（[^）]*(?:确认|未变|按此顺序)[^）]*）")




def _project_title(text: str) -> bool:
    return bool(re.search(r"\s*[｜|]\s*\S", text)) and not text.startswith(("•", "-", "·"))


def _parse_project_lines(lines: list[str], source_kind: str,
                         *, category: str = "unspecified",
                         locators: list[str] | None = None) -> list[dict[str, Any]]:
    projects: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    last_bullet = False
    for index, raw in enumerate(lines):
        line = " ".join(str(raw or "").split()).strip()
        if not line:
            continue
        if _project_title(line):
            if current:
                projects.append(current)
            parts = [x.strip() for x in re.split(r"\s*[｜|]\s*", line) if x.strip()]
            current = {
                "title": parts[0],
                "metadata": parts[1:],
                "bullets": [],
                "source_kind": source_kind,
                "category": category,
                "source_refs": [{"kind": source_kind,
                                 "locator": locators[index] if locators else f"line:{index}"}],
            }
            last_bullet = False
            continue
        if current is None:
            continue
        if line.startswith(("•", "-", "·")):
            current["bullets"].append(line.lstrip("•-· ").strip())
            last_bullet = True
            continue
        if re.fullmatch(r"20\d{2}\.\d{2}[–-](?:20\d{2}\.\d{2}|至今)", line) or line in {"在研", "独立产品设计 / AI 协同开发"}:
            current["metadata"].append(line)
            last_bullet = False
            continue
        if last_bullet and current["bullets"]:
            current["bullets"][-1] += " " + line
        else:
            current["metadata"].append(line)
    if current:
        projects.append(current)
    return projects


def _merge_projects(existing: list[dict[str, Any]], incoming: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out = [dict(x) for x in existing]
    def project_key(item: dict) -> str:
        # The same project has source-specific categories: MAX calls it
        # project_or_research while a resume may call it project or research.
        # Title plus date/institution/context distinguish same-title records.
        parts = [item.get("title", ""), *(item.get("metadata") or [])]
        return "|".join(re.sub(r"\s+", "", str(part)).casefold() for part in parts)

    def stable_id(item: dict) -> str:
        return "project-" + hashlib.sha256(project_key(item).encode()).hexdigest()[:16]

    for item in out:
        item.setdefault("id", stable_id(item))
    index = {project_key(item): i for i, item in enumerate(out)}
    for item in incoming:
        key = project_key(item)
        if not key:
            continue
        if key not in index:
            index[key] = len(out)
            out.append({**item, "id": stable_id(item)})
            continue
        old = out[index[key]]
        if old.get("category") in {None, "unspecified", "project_or_research"}:
            old["category"] = item.get("category") or old.get("category")
        elif (item.get("category") not in {None, "unspecified", "project_or_research", old.get("category")}):
            old["category"] = "ambiguous"
        old_meta = list(old.get("metadata") or [])
        for value in item.get("metadata") or []:
            if value not in old_meta:
                old_meta.append(value)
        old["metadata"] = old_meta
        old_bullets = list(old.get("bullets") or [])
        for value in item.get("bullets") or []:
            if value not in old_bullets:
                old_bullets.append(value)
        old["bullets"] = old_bullets
        sources = set(old.get("sources") or [old.get("source_kind")])
        sources.add(item.get("source_kind"))
        old["sources"] = sorted(x for x in sources if x)
        refs = list(old.get("source_refs") or [])
        for ref in item.get("source_refs") or []:
            if ref not in refs:
                refs.append(ref)
        old["source_refs"] = refs
    return out


def _education_records(profile: ApplicantProfile) -> list[dict[str, Any]]:
    """Structured education inventory derived only from sourced canonical fields."""
    records: list[dict[str, Any]] = []
    previous = {item.get("scope"): item.get("id")
                for item in (profile.collections.get("education_records") or [])
                if isinstance(item, dict) and isinstance(item.get("id"), str)}
    for scope in ("highest", "bachelor"):
        prefix = f"education.{scope}."
        fields = {key.removeprefix(prefix): field
                  for key, field in profile.fields.items() if key.startswith(prefix)}
        if not fields:
            continue
        identity = [fields.get(name).value if fields.get(name) else ""
                    for name in ("school", "degree", "major", "start_date", "graduation_date")]
        # Do not merge two partially described schools merely because one title overlaps.
        digest = hashlib.sha256(json.dumps([scope, identity], ensure_ascii=False,
                                           sort_keys=True, default=str).encode()).hexdigest()[:16]
        records.append({
            "id": previous.get(scope) or "education-" + digest,
            "scope": scope,
            "status": "PARTIAL" if any(value in (None, "") for value in identity) else "SOURCED",
            "fields": {name: {
                "value": field.value,
                "sources": [source.model_dump(mode="json") for source in field.sources],
                "last_verified": field.last_verified,
            } for name, field in fields.items()},
        })
    return records


def _resume_projects(text: str, source_kind: str) -> tuple[list[dict[str, Any]], str]:
    lines = [" ".join(line.split()).strip() for line in text.splitlines() if line.strip()]
    if not lines:
        return [], "EMPTY_SOURCE"
    groups: list[tuple[str, list[str]]] = []
    current: list[str] | None = None
    saw_section = False
    for line in lines:
        heading = re.sub(r"^[一二三四五六七八九十0-9]+[、.．]\s*", "", line).strip(" ：:")
        normalized_heading = heading.casefold()
        if normalized_heading in {"产品项目", "项目经历", "项目经验",
                                  "projects", "project experience", "product projects"}:
            current = []
            groups.append(("project", current))
            saw_section = True
            continue
        if normalized_heading in {"科研经历", "研究经历", "科研项目",
                                  "research", "research experience", "research projects"}:
            current = []
            groups.append(("research", current))
            saw_section = True
            continue
        if normalized_heading in {"技能与语言", "教育背景", "实习经历", "工作经历", "获奖荣誉",
                                  "skills", "education", "internships", "work experience",
                                  "awards"}:
            current = None
            continue
        if current is not None:
            current.append(line)
    records: list[dict[str, Any]] = []
    for category, group in groups:
        records.extend(_parse_project_lines(group, source_kind, category=category))
    return records, ("PARSED" if records else "UNPARSED_SECTION" if saw_section
                     else "NO_MATCHING_SECTION")

def file_sha256(path: str | Path) -> str:
    h = hashlib.sha256()
    with Path(path).open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _clean_value(key: str, raw: str):
    text = raw.strip()
    if any(marker in text for marker in PENDING_MARKERS):
        return None
    if key in BOOL_KEYS:
        if text.startswith(("是", "同意")):
            return True
        if text.startswith(("否", "不同意")):
            return False
        return None
    text = ANNOTATION_RE.sub("", text).strip()
    if key == "identity.driver_license":
        match = re.search(r"C[12]", text, re.I)
        return match.group(0).upper() if match else text
    if key == "certificates.mandarin":
        match = re.search(r"[一二三]级[甲乙]等", text)
        return match.group(0) if match else text
    if key == "certificates.computer":
        match = re.search(r"全国计算机等级考试二级（C语言程序设计）", text)
        return match.group(0) if match else text
    if key == "preferences.preferred_cities":
        return [x.strip().removesuffix("市") for x in re.split(r"[>＞]", text) if x.strip()]
    if key == "education.highest.graduation_date":
        m = re.match(r"(20\d{2}-\d{2}(?:-\d{2})?)", text)
        if m:
            return m.group(1)
    return text


def _evidence_priority(kind: str, user_confirmed: bool = False) -> int:
    if kind == "user_explicit":
        return 500
    if kind in {"max_docx", "resume_pdf"}:
        return 420 if user_confirmed else 400
    if kind == "verified_history":
        return 300
    if kind == "legacy_profile":
        return 200
    return 100


def _field_priority(field: ProfileField) -> int:
    priorities = [_evidence_priority(source.kind, field.user_confirmed) for source in field.sources]
    return max(priorities, default=0)


class ProfileBuilder:
    def __init__(self):
        self.profile = ApplicantProfile()
        self.conflicts: list[dict[str, Any]] = []

    def add_field(self, key: str, value: Any, source: EvidenceRef, *, confidence: float = 0.98,
                  user_confirmed: bool = False, last_verified: str | None = None,
                  normalization: dict[str, Any] | None = None):
        if value in (None, ""):
            return
        existing = self.profile.fields.get(key)
        incoming = ProfileField(
            value=value,
            sources=[source],
            confidence=confidence,
            last_verified=last_verified,
            aliases=list(aliases_for(key)),
            normalization=normalization or {},
            user_confirmed=user_confirmed,
            sensitive=is_sensitive_key(key),
        )
        if existing is None:
            self.profile.fields[key] = incoming
            return
        if existing.value == value:
            existing.sources.append(source)
            existing.confidence = max(existing.confidence, confidence)
            existing.user_confirmed = existing.user_confirmed or user_confirmed
            existing.last_verified = existing.last_verified or last_verified
            return
        existing_rank = (_field_priority(existing), existing.confidence)
        incoming_rank = (
            _evidence_priority(source.kind, incoming.user_confirmed),
            incoming.confidence,
        )
        self.conflicts.append({
            "key": key,
            "existing_value": existing.value,
            "incoming_value": value,
            "existing_sources": [s.model_dump() for s in existing.sources],
            "incoming_source": source.model_dump(),
        })
        if incoming_rank > existing_rank:
            self.profile.fields[key] = incoming

    def import_max_docx(self, path: str | Path):
        path = Path(path).expanduser().resolve()
        doc = Document(path)
        source_doc = EvidenceRef(kind="max_docx", path=str(path), note=f"sha256:{file_sha256(path)}")
        self.profile.source_documents.append(source_doc)
        project_lines: list[dict[str, Any]] = []
        in_projects = False
        for index, paragraph in enumerate(doc.paragraphs):
            text = " ".join((paragraph.text or "").split()).strip()
            if not text:
                continue
            if text.startswith("四、科研与项目经历"):
                in_projects = True
                continue
            if text.startswith("五、技能、语言与证书"):
                in_projects = False
            if in_projects:
                project_lines.append({"text": text, "source": f"paragraph:{index}"})
            if "：" not in text:
                continue
            label, raw = text.split("：", 1)
            key = LABEL_MAP.get(label.strip())
            if not key:
                continue
            value = _clean_value(key, raw)
            verified = CONFIRM_RE.search(text)
            verified_at = verified.group(1) if verified else None
            self.add_field(
                key,
                value,
                EvidenceRef(kind="max_docx", path=str(path), locator=f"paragraph:{index}", verified_at=verified_at),
                confidence=1.0 if verified else 0.98,
                user_confirmed=bool(verified),
                last_verified=verified_at,
                normalization={"strategy": "max_label_map", "source_label": label.strip()},
            )
            if label.strip() == "本科专业排名":
                pass

        for index, paragraph in enumerate(doc.paragraphs):
            text = " ".join((paragraph.text or "").split()).strip()
            if text.startswith("本科专业排名："):
                raw = text.split("：", 1)[1]
                rank = raw.split("；", 1)[0].strip()
                band_match = re.search(r"排名档位：([^；]+)$", raw)
                src = EvidenceRef(kind="max_docx", path=str(path), locator=f"paragraph:{index}")
                self.add_field("education.bachelor.rank", rank, src)
                if band_match:
                    self.add_field("education.bachelor.rank_band", band_match.group(1).strip(), src)

        self.profile.collections["projects_from_max"] = project_lines
        parsed_max = _parse_project_lines([x["text"] for x in project_lines], "max_docx",
                                          category="project_or_research",
                                          locators=[x["source"] for x in project_lines])
        self.profile.collections["max_project_parse_status"] = (
            "PARSED" if parsed_max else "UNPARSED_SECTION" if project_lines else "EMPTY_SECTION"
        )
        self.profile.collections["projects"] = _merge_projects(
            self.profile.collections.get("projects") or [],
            parsed_max,
        )
        return self

    def import_legacy_profile(self, path: str | Path):
        path = Path(path).expanduser().resolve()
        if not path.is_file():
            return self
        data = json.loads(path.read_text(encoding="utf-8"))
        mappings = {
            "identity.full_name": ("identity", "full_name"),
            "identity.phone": ("identity", "phone"),
            "identity.email": ("identity", "email"),
            "identity.current_residence": ("identity", "location"),
            "identity.current_city": ("identity", "current_city"),
            "identity.pre_gaokao_domicile": ("identity", "household_before_gaokao"),
            "identity.domicile": ("identity", "household_address"),
            "identity.family_address": ("identity", "family_address"),
            "identity.student_origin": ("identity", "student_origin"),
            "identity.household_type": ("identity", "household_type"),
            "identity.health_status": ("identity", "health_status"),
            "identity.personnel_file_place": ("identity", "personnel_file_place"),
            "identity.foreign_residency_status": ("identity", "foreign_residency_status"),
            "compliance.coamc_employee_recusal_requirements_met": ("compliance", "coamc_employee_recusal_requirements_met"),
            "family.primary.name": ("family", "primary", "name"),
            "family.primary.relationship": ("family", "primary", "relationship"),
            "family.primary.work_unit": ("family", "primary", "work_unit"),
            "family.primary.department_title": ("family", "primary", "department_title"),
            "family.primary.work_location": ("family", "primary", "work_location"),
            "policy.auto_accept_privacy_terms": ("application_policy", "auto_accept_privacy_terms"),
            "policy.auto_accept_truth_submission_declarations": ("application_policy", "auto_accept_truth_submission_declarations"),
            "policy.auto_decide_company_legal_compliance": ("application_policy", "auto_decide_company_legal_compliance"),
            "policy.final_submission_requires_user_click": ("application_policy", "final_submission_requires_user_click"),
            "education.highest.school": ("education", "school"),
            "education.highest.degree": ("education", "degree"),
            "education.highest.major": ("education", "major"),
            "education.highest.graduation_date": ("education", "graduation_date"),
        }
        for key, route in mappings.items():
            cur: Any = data
            for part in route:
                cur = cur.get(part) if isinstance(cur, dict) else None
            if cur not in (None, ""):
                self.add_field(
                    key, cur,
                    EvidenceRef(kind="legacy_profile", path=str(path), locator=".".join(route)),
                    confidence=0.90,
                    normalization={"strategy": "legacy_mapping", "source_key": ".".join(route)},
                )
        return self

    def import_resume(self, path: str | Path):
        path = Path(path).expanduser().resolve()
        if not path.is_file():
            return self
        if path.suffix.lower() == ".pdf":
            return self.import_resume_pdf(path)
        if path.suffix.lower() == ".docx":
            doc = Document(path)
            text = "\n".join((paragraph.text or "") for paragraph in doc.paragraphs)
            kind = "resume_docx"
            self.profile.source_documents.append(
                EvidenceRef(kind=kind, path=str(path), note=f"sha256:{file_sha256(path)}")
            )
            if text.strip():
                self.profile.collections["resume_text_snapshot"] = text.strip()
                parsed_resume, parse_status = _resume_projects(text, kind)
                self.profile.collections["resume_project_parse_status"] = parse_status
                self.profile.collections["projects"] = _merge_projects(
                    self.profile.collections.get("projects") or [],
                    parsed_resume,
                )
            else:
                self.profile.collections["resume_project_parse_status"] = "EMPTY_SOURCE"
            self.profile.assets["resume"] = AssetRef(
                path=str(path), kind=kind, sha256=file_sha256(path),
                source=EvidenceRef(kind=kind, path=str(path)),
            )
            return self
        return self.add_asset("resume", path, f"resume_{path.suffix.lower().lstrip('.') or 'file'}")

    def import_resume_pdf(self, path: str | Path):
        path = Path(path).expanduser().resolve()
        if not path.is_file():
            return self
        text = ""
        try:
            import pymupdf
            with pymupdf.open(path) as pdf:
                text = "\n".join(page.get_text("text") for page in pdf)
        except Exception as exc:
            self.profile.collections["resume_parse_warning"] = type(exc).__name__
        self.profile.source_documents.append(
            EvidenceRef(kind="resume_pdf", path=str(path), note=f"sha256:{file_sha256(path)}")
        )
        if text.strip():
            self.profile.collections["resume_text_snapshot"] = text.strip()
            parsed_resume, parse_status = _resume_projects(text, "resume_pdf")
            self.profile.collections["resume_project_parse_status"] = parse_status
            self.profile.collections["projects"] = _merge_projects(
                self.profile.collections.get("projects") or [],
                parsed_resume,
            )
        else:
            self.profile.collections["resume_project_parse_status"] = (
                "PARSE_FAILED" if self.profile.collections.get("resume_parse_warning")
                else "EMPTY_SOURCE"
            )
        self.profile.assets["resume"] = AssetRef(
            path=str(path), kind="resume_pdf", sha256=file_sha256(path),
            source=EvidenceRef(kind="resume_pdf", path=str(path)),
        )
        return self

    def add_asset(self, name: str, path: str | Path, kind: str):
        path = Path(path).expanduser().resolve()
        if path.is_file():
            self.profile.assets[name] = AssetRef(
                path=str(path), kind=kind, sha256=file_sha256(path),
                source=EvidenceRef(kind="local_asset", path=str(path)),
            )
        return self

    def import_historical_json(self, scope: str, path: str | Path, *, reconfirm_keys: list[str] | None = None):
        path = Path(path).expanduser().resolve()
        if not path.is_file():
            return self
        data = json.loads(path.read_text(encoding="utf-8"))
        records = self.profile.collections.setdefault("historical_answers", {})
        records[scope] = {
            "source": {"kind": "historical_application", "path": str(path), "imported_at": _now()},
            "answers": data,
            "requires_reconfirmation": list(reconfirm_keys or []),
        }

        canonical = data.get("canonical_fields") if isinstance(data, dict) else None
        if isinstance(canonical, dict):
            for key, raw in canonical.items():
                item = raw if isinstance(raw, dict) else {"value": raw}
                if item.get("value") in (None, ""):
                    continue
                self.add_field(
                    str(key),
                    item.get("value"),
                    EvidenceRef(
                        kind="verified_history",
                        path=str(path),
                        locator=f"canonical_fields.{key}",
                        verified_at=item.get("verified_at"),
                        note=f"scope:{scope}",
                    ),
                    confidence=min(float(item.get("confidence", 0.92)), 0.95),
                    user_confirmed=bool(item.get("user_confirmed", False)),
                    last_verified=item.get("verified_at"),
                    normalization={"strategy": "verified_history", "scope": scope},
                )
        return self

    def import_user_overrides(self, path: str | Path):
        path = Path(path).expanduser().resolve()
        if not path.is_file():
            return self
        raw = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(raw, dict) or (raw and "fields" not in raw):
            raise ValueError("existing output is not a canonical profile")
        existing = ApplicantProfile.model_validate(raw)
        for collection in ("profile_revisions", "profile_write_version", "fact_exclusions"):
            if collection in existing.collections:
                self.profile.collections[collection] = existing.collections[collection]
        for key, field in existing.fields.items():
            if not any(source.kind == "user_explicit" for source in field.sources):
                continue
            self.add_field(
                key,
                field.value,
                EvidenceRef(
                    kind="user_explicit",
                    path=str(path),
                    locator=key,
                    verified_at=field.last_verified,
                    note="preserved canonical user override",
                ),
                confidence=1.0,
                user_confirmed=True,
                last_verified=field.last_verified,
                normalization={"strategy": "user_explicit"},
            )
        return self

    def build(self) -> ApplicantProfile:
        if self.conflicts:
            self.profile.collections["conflicts"] = self.conflicts
        self.profile.collections["education_records"] = _education_records(self.profile)
        return self.profile


def set_user_confirmed_field(
    profile_path: str | Path,
    key: str,
    value: Any,
    *,
    note: str | None = None,
) -> Path:
    profile_path = Path(profile_path).expanduser().resolve()
    with _profile_lock(profile_path):
        return _set_user_confirmed_field_locked(profile_path, key, value, note=note)


def set_project_exclusion(profile_path: str | Path, task_id: str,
                          project_id: str, reason: str) -> Path:
    """Record an explicit per-application omission without deleting inventory."""
    if not re.fullmatch(r"[A-Za-z0-9_.:-]{8,120}", task_id):
        raise ValueError("invalid task scope")
    if not isinstance(reason, str) or not reason.strip() or len(reason) > 1000:
        raise ValueError("explicit exclusion reason required")
    path = Path(profile_path).expanduser().resolve()
    with _profile_lock(path):
        profile = ApplicantProfile.model_validate_json(path.read_text(encoding="utf-8"))
        projects = profile.collections.get("projects") or []
        if not any(isinstance(item, dict) and item.get("id") == project_id for item in projects):
            raise ValueError("project identity not in canonical inventory")
        scopes = profile.collections.setdefault("fact_exclusions", {})
        decisions = scopes.setdefault(task_id, {})
        decisions[project_id] = {
            "reason": reason.strip(), "source": "user_explicit_task", "at": _now(),
        }
        return _write_profile_locked(profile, path)


def projects_for_task(profile: ApplicantProfile, task_id: str) -> list[dict[str, Any]]:
    excluded = (profile.collections.get("fact_exclusions") or {}).get(task_id, {})
    return [item for item in (profile.collections.get("projects") or [])
            if item.get("id") not in excluded]


def _set_user_confirmed_field_locked(profile_path: Path, key: str, value: Any,
                                     *, note: str | None) -> Path:
    original = json.loads(profile_path.read_text(encoding="utf-8"))
    if not isinstance(original, dict) or (original and "fields" not in original):
        raise ValueError("reusable fact requires a canonical profile")
    profile = ApplicantProfile.model_validate(original)
    if note and any(item.get("note") == note for item in
                    (profile.collections.get("profile_revisions") or [])
                    if isinstance(item, dict)):
        return profile_path
    previous = profile.fields.get(key)
    verified_at = _now()
    revisions = profile.collections.setdefault("profile_revisions", [])
    revisions.append({
        "key": key,
        "at": verified_at,
        "previous_source_kinds": [x.kind for x in previous.sources] if previous else [],
        "previous_value": previous.value if previous else None,
        "confirmed_value": value,
        "source": "user_explicit",
        "note": note,
    })
    sources = list(previous.sources) if previous else []
    sources.append(EvidenceRef(
        kind="user_explicit",
        path=str(profile_path),
        locator=key,
        verified_at=verified_at,
        note=note,
    ))
    profile.fields[key] = ProfileField(
        value=value,
        sources=sources,
        confidence=1.0,
        last_verified=verified_at,
        aliases=list(aliases_for(key)),
        normalization={"strategy": "user_explicit"},
        user_confirmed=True,
        sensitive=is_sensitive_key(key),
    )
    if key.startswith("education."):
        profile.collections["education_records"] = _education_records(profile)
    return _write_profile_locked(profile, profile_path)


@contextmanager
def _profile_lock(path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    lock_path = path.with_name(path.name + ".lock")
    with lock_path.open("a+") as lock:
        lock_path.chmod(0o600)
        fcntl.flock(lock, fcntl.LOCK_EX)
        try:
            yield
        finally:
            fcntl.flock(lock, fcntl.LOCK_UN)


def write_profile(profile: ApplicantProfile, path: str | Path) -> Path:
    path = Path(path).expanduser().resolve()
    path.parent.mkdir(parents=True, exist_ok=True)
    with _profile_lock(path):
        return _write_profile_locked(profile, path)


def _write_profile_locked(profile: ApplicantProfile, path: Path) -> Path:
    if path.exists():
        # An unreadable canonical profile must never be replaced by an import.
        raw = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(raw, dict) or (raw and "fields" not in raw):
            raise ValueError("existing output is not a canonical profile")
        on_disk = ApplicantProfile.model_validate(raw)
        current_version = int(on_disk.collections.get("profile_write_version", 0))
        incoming_version = int(profile.collections.get("profile_write_version", 0))
        if incoming_version != current_version:
            raise RuntimeError("stale canonical profile version")
        path.chmod(0o600)
        snapshot = path.with_name(path.name + ".pre-jcr04")
        try:
            os.link(path, snapshot)
        except FileExistsError:
            pass
    updated = profile.model_copy(deep=True)
    updated.collections["profile_write_version"] = (
        int(updated.collections.get("profile_write_version", 0)) + 1
    )
    payload = updated.model_dump_json(indent=2).encode("utf-8")
    fd, name = tempfile.mkstemp(dir=path.parent, prefix=".profile-write-")
    temporary = Path(name)
    try:
        os.fchmod(fd, 0o600)
        with os.fdopen(fd, "wb") as file:
            file.write(payload)
            file.flush()
            os.fsync(file.fileno())
        os.replace(temporary, path)
        profile.collections["profile_write_version"] = updated.collections["profile_write_version"]
        directory = os.open(path.parent, os.O_RDONLY)
        try:
            os.fsync(directory)
        finally:
            os.close(directory)
    finally:
        temporary.unlink(missing_ok=True)
    return path
